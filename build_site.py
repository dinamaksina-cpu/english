from openpyxl import load_workbook
from pathlib import Path
import json, re, html

src=Path('/mnt/data/Oxford_3000_UA.xlsx')
out=Path('/mnt/data/english5000_pwa')
wb=load_workbook(src, read_only=True, data_only=True)
ws=wb['Oxford 3000']
rows=[]
for r in ws.iter_rows(min_row=2, values_only=True):
    idx, word, tr, pron, ipa = r
    rows.append({'id':int(idx),'word':str(word),'translation':str(tr or ''),'pron':str(pron or ''),'ipa':str(ipa or '')})

TOPICS={
'Food & Drink': set('apple banana bread breakfast butter cake carrot cheese chicken chocolate coffee cook cooking cream cup diet dinner dish drink egg fish food fruit hungry ice kitchen lunch meal meat milk orange potato restaurant rice salad salt sandwich sugar tea tomato vegetable water wine'.split()),
'Home': set('apartment bath bathroom bed bedroom chair clean cleaning cooker desk door flat floor furniture garden home house key kitchen lamp living room mirror room shelf shower sofa table toilet wall window'.split()),
'Work & Business': set('business career company customer job manager meeting office project report salary sell service staff team work worker'.split()),
'Education': set('book class classroom college course dictionary education exam homework learn lesson library school student study teacher test university'.split()),
'Travel': set('airport beach booking border camp camping country flight holiday hotel journey map passport reservation suitcase tourism tourist tour travel trip visit'.split()),
'Transport': set('bicycle bike boat bus car drive driver driving flight motorbike plane railway road station taxi train transport truck'.split()),
'Health & Body': set('arm back blood body brain doctor ear eye face finger foot hair hand health heart hospital knee leg medicine neck nose nurse pain sick skin stomach tooth teeth'.split()),
'Family & People': set('adult aunt baby boy boyfriend brother child cousin dad daughter family father friend girl girlfriend grandfather grandmother husband man mother mum parent person sister son uncle wife woman'.split()),
'Emotions & Personality': set('afraid angry bored boring calm confident excited friendly funny happy kind lonely nervous nice proud sad serious shy surprised tired worried'.split()),
'Nature & Weather': set('air animal beach bird cloud cold earth environment farm field fire flower forest garden grass hot lake mountain nature ocean plant rain river sea sky snow star sun tree warm weather wind'.split()),
'Technology': set('app blog camera cd computer digital dvd email file internet laptop message mobile online password phone photo screen software technology tv video website'.split()),
'Money & Shopping': set('bank bill buy card cash cheap cost customer dollar euro market money pay price product sale sell shop shopping store'.split()),
'Time & Dates': set('afternoon april august century date day december evening february friday future hour january july june march may minute monday month morning night november october saturday second september sunday thursday time today tomorrow tuesday wednesday week weekend year yesterday'.split()),
'Clothes & Appearance': set('bag belt blonde boot brown clothes coat dress fashion hat jacket jeans shoe shoes skirt suit sweater trousers wear'.split()),
'Sports & Fitness': set('ball basketball bike climb club football game gym play player race run running ski sport swim swimming team tennis'.split()),
'Culture & Media': set('art artist article book cinema concert culture dance dancer film media movie music news newspaper painting photo radio song story theatre tv video'.split()),
'Animals': set('animal bird cat chicken cow dog fish horse insect pet sheep'.split()),
'Places & City': set('area bank beach building cafe centre city club college country hospital hotel library market office park place restaurant road school shop station street university'.split()),
'Communication & Language': set('answer ask call conversation describe description email explain language letter listen message name question read say speak talk tell text translate translation word write'.split()),
}
# phrases need exact handling. Keep topic labels broad and pedagogical, not official Oxford topics.
def norm(s): return re.sub(r"[^a-z0-9' ]+",'',s.lower()).strip()
for item in rows:
    w=norm(item['word'])
    cats=[]
    for t, words in TOPICS.items():
        if w in words:
            cats.append(t)
    item['topics']=cats or ['General']

(out/'words.json').write_text(json.dumps(rows,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
print('rows',len(rows),'topics',len(TOPICS)+1)
